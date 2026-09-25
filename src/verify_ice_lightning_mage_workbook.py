#!/usr/bin/env python3
"""
Cross-checks Ice-Lightning-Mage-DPS-Calculator.xlsx against an independent Python port of the
same logic, using the workbook's own default Inputs values. Then loads the live workbook with
the `formulas` package (a real formula evaluator, not just the static values openpyxl wrote) and
diffs Total DPS + every per-skill DPS cell (Calc!O<row>) against this script's independent
numbers. Mirrors the sibling verify_fp_mage_workbook.py / verify_night_lord_workbook.py scripts
(read-only references, not modified) in structure and method.
"""
import json
import math
import re
import sys
from pathlib import Path

import numpy as np

REPO = Path(__file__).resolve().parent.parent
XLSX_PATH = REPO / "Ice-Lightning-Mage" / "Ice-Lightning-Mage-DPS-Calculator.xlsx"
sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_ice_lightning_mage_workbook import (  # noqa: E402
    ROW, IN, UNLOCK_LEVEL, SUMMARY_ROW, CONTENT_TYPES, PER_CONTENT_TYPE_INPUT_KEYS,
)

FACTOR_TABLE = json.loads((REPO / "data/factor_table.json").read_text())
FACTOR_TABLE = {int(k): v for k, v in FACTOR_TABLE.items()}


def get_factor(level, idx):
    level = min(300, max(1, round(level)))
    return FACTOR_TABLE[level][idx]


def input_level(job_step, level, skill1, skill2, skill3, skill4, skill_all):
    if job_step == 1:
        return 60 + skill1 + skill_all
    if job_step == 2:
        return 90 + skill2 + skill_all
    if job_step == 3:
        return 120 + skill3 + skill_all
    return max(0, (level - 100) * 3) + skill4 + skill_all


import openpyxl  # noqa: E402

_inputs_ws = openpyxl.load_workbook(XLSX_PATH)["Inputs"]

# Inputs is now per-content-type (columns C-L, one per CONTENT_TYPES entry, resolved into column
# B via an INDEX/MATCH formula keyed on the active Content Type) — read the RAW per-content-type
# cell for the currently active content type directly, since column B itself now holds a formula
# string (not a static value) that plain openpyxl can't evaluate. Mirrors the same fix applied to
# verify_fp_mage_workbook.py / verify_bishop_workbook.py's own _in().
_col_for_ct = {}
for _c in range(3, 3 + len(CONTENT_TYPES)):
    _name = _inputs_ws.cell(row=2, column=_c).value
    if _name in CONTENT_TYPES:
        _col_for_ct[_name] = _c
_active_content_type = _inputs_ws.cell(row=IN["content_type"], column=2).value


def _in(key):
    if key not in PER_CONTENT_TYPE_INPUT_KEYS:
        return _inputs_ws.cell(row=IN[key], column=2).value
    col = _col_for_ct.get(_active_content_type)
    value = _inputs_ws.cell(row=IN[key], column=col).value if col else None
    # A handful of per-content-type rows (Boss/Normal Emphasis, Max Enemies Actually In Range)
    # are blank/gray for content types where they're not applicable (see
    # INPUT_ROW_APPLICABLE_CONTENT_TYPES in build_inputs_sheet) — guard against None so downstream
    # numeric use (e.g. min(targets, max_enemies_hit)) doesn't blow up on a blank cell.
    return value if value is not None else 0


level = _in("level")
def _compute_monster_type(content_type):
    if content_type == "PvP":
        return "pvp"
    if content_type in ("Breakthrough", "Hero Dungeon"):
        return "breakthrough"
    if content_type in ("EXP Dungeon", "Equipment Dungeon", "Chapter Hunt"):
        return "normal"
    return "boss"


def _compute_breakthrough_stage_index(chapter, stage):
    if chapter == 28:
        return stage - 9
    return stage + 14 * max(0, min(chapter - 1, 38) - 28) + 19 * max(0, chapter - 39)


def _compute_monster_defense(content_type, chapter, stage, defense):
    if content_type == "Chapter Boss":
        return 3200 + 50 * (chapter - 28)
    if content_type == "World Boss":
        return 62100
    if content_type == "Weapon Dungeon":
        return stage * 50
    if content_type == "Enhancement Dungeon":
        return 950 + stage * 50
    if content_type in ("EXP Dungeon", "Equipment Dungeon"):
        return 250 + stage * 50
    if content_type == "Hero Dungeon":
        return 650 + stage * 50
    if content_type in ("Breakthrough", "Chapter Hunt"):
        return 4860 + 20 * _compute_breakthrough_stage_index(chapter, stage)
    return defense  # PvP fallback — manual, opponent-dependent


def _compute_fight_duration(content_type):
    return {
        "Chapter Hunt": 0, "Weapon Dungeon": 22, "Equipment Dungeon": 40,
        "Enhancement Dungeon": 25, "Hero Dungeon": 50, "EXP Dungeon": 40,
        "World Boss": 75, "Chapter Boss": 70, "Breakthrough": 40,
    }.get(content_type, 0)


def _parse_chapter_stage(raw):
    raw = str(raw)
    if "-" in raw:
        left, right = raw.split("-", 1)
        try:
            return float(left), float(right)
        except ValueError:
            return 28.0, 9.0
    try:
        v = float(raw)
        return v, v
    except ValueError:
        return 28.0, 9.0


def _compute_breakthrough_normal_weight_pct(choice):
    return {
        "More Normal": 70, "A Little More Normal": 60, "Equal": 50,
        "A Little More Boss": 40, "More Boss": 30,
    }.get(choice, 60)


content_type = _in("content_type")
chapter, stage = _parse_chapter_stage(_in("chapter_stage"))
defense = _in("defense")
monster_type = _compute_monster_type(content_type)
breakthrough_normal_weight_pct = _compute_breakthrough_normal_weight_pct(_in("boss_normal_weight_choice"))
normal_weight = (
    1.0 if monster_type == "normal"
    else breakthrough_normal_weight_pct / 100 if monster_type == "breakthrough"
    else 0.0
)
skill1, skill2, skill3, skill4, skill_all = (
    _in("skill_lvl_1st"), _in("skill_lvl_2nd"), _in("skill_lvl_3rd"), _in("skill_lvl_4th"), _in("skill_lvl_all")
)
monster_defense = _compute_monster_defense(content_type, chapter, stage, defense)
flat_attack, attack_pct = _in("flat_attack"), _in("attack_pct")
attack = flat_attack * (1 + attack_pct / 100)
flat_int, int_pct, luk = _in("flat_int"), _in("int_pct"), _in("luk")
stat_damage = (flat_int * (1 + int_pct / 100)) * 0.01 + luk * 0.0025
damage = _in("damage")
damage_amp = _in("damage_amp")
boss_damage = _in("boss_damage")
normal_damage = _in("normal_damage")
final_damage = _in("final_damage")
min_damage = _in("min_damage")
max_damage = _in("max_damage")
crit_rate = _in("crit_rate")
crit_damage = _in("crit_damage")
attack_speed_base = _in("attack_speed")
def_pen = _in("def_pen")
basic_attack_damage = _in("basic_attack_damage")
skill_damage = _in("skill_damage")
skill_cooldown_decrease = _in("skill_cooldown_decrease")
basic_attack_target_increase = _in("basic_attack_target_increase")
buff_duration_increase_pct = _in("buff_duration_increase_pct")
fight_duration = _compute_fight_duration(content_type)
max_enemies_hit = _in("max_enemies_hit")

fixed_duration_active = monster_type != "pvp" and fight_duration > 0

basic_lvl = input_level(4, level, skill1, skill2, skill3, skill4, skill_all)
skill_coefficient_base = 290 * get_factor(basic_lvl, 21) / 1000
print(f"SKILL_COEFFICIENT (base) = {skill_coefficient_base:.4f}")


def level_gated_sum(pairs):
    return sum(inc for lvl, inc in pairs.items() if level >= lvl)


# Chain Lightning's own full mastery chain (folded into its own row, not separate rows).
chain_lightning_mastery = level_gated_sum({102: 10, 106: 1, 116: 1, 120: 1, 128: 1, 132: 1})
chain_lightning_boss_mastery = level_gated_sum({111: 10, 124: 10})
chain_lightning_hits = 6 if level >= 136 else 5

thunder_sphere_mastery_normal = 150 + (100 if level >= 94 else 0)
thunder_sphere_icd = (2 * 0.7) if level >= 104 else 2
thunder_sphere_cooldown = (30 * 0.7) if level >= 78 else 30

glacier_wall_cooldown = (22 * 0.7) if level >= 90 else 22
magic_guard_cooldown = (30 * 0.7) if level >= 21 else 30
meditation_duration = (15 * 1.3) if level >= 44 else 15

frozen_orb_icd = (0.5 * 0.5) if level >= 134 else 0.5
elquines_icd = (4 * 0.8) if level >= 126 else 4
elquines_targets = 6 if level >= 138 else 3

SKILLS = {
    "CHAIN_LIGHTNING": dict(job_step=4, cooldown=None, hits=chain_lightning_hits, base=None, fidx=None, scales=True,
                             mastery=chain_lightning_mastery, mastery_boss=chain_lightning_boss_mastery, mastery_normal=0,
                             costs_action=False),
    "THUNDER_BOLT": dict(job_step=2, cooldown=18, hits=3, base=1800, fidx=12, scales=True,
                          mastery=level_gated_sum({39: 70}), mastery_boss=0, mastery_normal=0, costs_action=True),
    "GLACIER_WALL": dict(job_step=3, cooldown=glacier_wall_cooldown, hits=3, base=2900, fidx=12, scales=True,
                          mastery=level_gated_sum({68: 80}), mastery_boss=0, mastery_normal=0, costs_action=True),
    "THUNDER_SPHERE": dict(job_step=3, cooldown=thunder_sphere_cooldown, hits=3, icd=thunder_sphere_icd, window=10,
                            base=1000, fidx=12, scales=True, mastery=level_gated_sum({73: 50}), mastery_boss=0,
                            mastery_normal=thunder_sphere_mastery_normal, costs_action=True),
    "FREEZING_BREATH": dict(job_step=4, cooldown=40, hits=1, icd=0.5, window=5, base=8000, fidx=12, scales=True,
                             mastery=level_gated_sum({108: 50}), mastery_boss=0, mastery_normal=0, costs_action=True),
    "BLIZZARD": dict(job_step=4, cooldown=33, hits=3, base=6000, fidx=12, scales=True,
                      mastery=level_gated_sum({130: 40}), mastery_boss=0, mastery_normal=0, costs_action=True),
    "BLIZZARD_FINAL_ATTACK": dict(job_step=4, cooldown=33, hits=1, base=9500, fidx=21, scales=True, chance=30, rolls=3,
                                   mastery=0, mastery_boss=0, mastery_normal=0, costs_action=False),
    "FROZEN_ORB": dict(job_step=4, cooldown=23, hits=1, icd=frozen_orb_icd, window=6, base=9000, fidx=12, scales=True,
                        mastery=level_gated_sum({122: 50}), mastery_boss=0, mastery_normal=0, costs_action=True),
    "ELQUINES": dict(job_step=4, cooldown=80, hits=1, icd=elquines_icd, window=30, base=35000, fidx=12, scales=True,
                      mastery=level_gated_sum({138: 50}), mastery_boss=0, mastery_normal=0, costs_action=True),
}
BUFFS = {
    "MAGIC_GUARD": dict(job_step=1, cooldown=magic_guard_cooldown, duration=15, base=120, fidx=21, scales=True,
                         target="ATTACK", costs_action=True),
    "MEDITATION": dict(job_step=2, cooldown=30, duration=meditation_duration, base=200, fidx=21, scales=True,
                        target="ATTACK", costs_action=True),
    "INFINITY": dict(job_step=4, cooldown=30, duration=15, base=150, fidx=21, scales=True,
                      target="FINAL_DAMAGE", costs_action=True, ramp_mult=64 / 45),
}
PASSIVES = {
    "BUFF_MASTERY": dict(job_step=4, base=100, fidx=22, scales=True),
    "ARCANE_AIM": dict(job_step=4, base=30, fidx=22, scales=True),
    "ELEMENT_AMPLIFICATION": dict(job_step=3, base=150, fidx=22, scales=True),
    "ELEMENTAL_RESET": dict(job_step=3, base=120, fidx=22, scales=True),
    "MP_EATER_MP_BOOST": dict(job_step=2, base=70, fidx=0, scales=False),
}
MAPLE_TARGETS = {"THUNDER_BOLT": 1000, "GLACIER_WALL": 300, "THUNDER_SPHERE": 200}
NORMAL_MONSTER_TARGETS = {
    "THUNDER_BOLT": 8, "GLACIER_WALL": 8, "THUNDER_SPHERE": 6, "FREEZING_BREATH": 10,
    "BLIZZARD": 7, "BLIZZARD_FINAL_ATTACK": 1, "FROZEN_ORB": 10, "ELQUINES": elquines_targets,
}


def boss_normal_multiplier(key, s):
    """Correctly blends Boss/Normal Monster Damage% and target count: each branch gets its own
    full (1+damage%/100)*targets treatment. `s` is the skill's own mastery dict."""
    if monster_type == "pvp":
        return 1
    if key == "CHAIN_LIGHTNING":
        targets = 6 + basic_attack_target_increase
    else:
        targets = NORMAL_MONSTER_TARGETS.get(key, 1)
    targets = min(targets, max_enemies_hit)
    boss_term = boss_damage + s.get("mastery_boss", 0) + monster_dmg_bonus
    normal_term = normal_damage + s.get("mastery_normal", 0) + monster_dmg_bonus
    boss_branch = 1 + boss_term / 100
    normal_branch = (1 + normal_term / 100) * targets
    return (1 - normal_weight) * boss_branch + normal_weight * normal_branch


def unlocked(key):
    threshold = UNLOCK_LEVEL.get(key)
    return threshold is None or level >= threshold


PVP_FIGHT_DURATION = 15


def eff_cooldown(cooldown, costs_action):
    if monster_type == "pvp":
        return PVP_FIGHT_DURATION
    return max(0.1, cooldown - (skill_cooldown_decrease if costs_action else 0))


def eff_duration(duration):
    return duration * (1 + (buff_duration_increase_pct + buff_mastery_pct) / 100)


def exact_casts(cooldown, duration=None):
    d = fight_duration if duration is None else duration
    return math.floor(d / cooldown) + 1


def exact_total_hits(cooldown, hits_per_cast, icd, window, duration=None):
    d = fight_duration if duration is None else duration
    casts = exact_casts(cooldown, d)
    remaining_after_last = max(0.0, d - (casts - 1) * cooldown)
    if icd:
        full_window_ticks = window / icd
        last_cast_ticks = min(window, remaining_after_last) / icd
    else:
        full_window_ticks = 1
        last_cast_ticks = 1
    return hits_per_cast * ((casts - 1) * full_window_ticks + last_cast_ticks)


def exact_buff_uptime(cooldown, buff_duration):
    casts = exact_casts(cooldown)
    remaining_after_last = max(0.0, fight_duration - (casts - 1) * cooldown)
    last_uptime = min(buff_duration, remaining_after_last)
    return (casts - 1) * buff_duration + last_uptime


def coeff_pct(base, fidx, scales, job_step):
    lvl = input_level(job_step, level, skill1, skill2, skill3, skill4, skill_all)
    if not scales:
        return base / 10
    return (base / 10) * (get_factor(lvl, fidx) / 1000)


def passive_pct(key):
    p = PASSIVES[key]
    return coeff_pct(p["base"], p["fidx"], p["scales"], p["job_step"]) if unlocked(key) else 0.0


buff_mastery_pct = passive_pct("BUFF_MASTERY")
arcane_aim_pct = passive_pct("ARCANE_AIM")
element_amp_pct = passive_pct("ELEMENT_AMPLIFICATION")
elemental_reset_pct = passive_pct("ELEMENTAL_RESET")
mp_eater_as_pct = passive_pct("MP_EATER_MP_BOOST")

# Freezing Breath - Weaken (Mastery Lv.118): +15% Damage Taken debuff, 30s, duty-cycle averaged
# off Freezing Breath's own (patched) 40s cooldown — steady-state only (no fixed-duration
# exactness, a documented simplification for this secondary mastery effect). Elemental Reset's
# own Damage Taken bonus is the SAME kind of bonus (target damage-taken%), so it sums into this
# same bucket instead of being a separate multiplicative factor.
if level >= 118:
    fb_eff_cd = eff_cooldown(40, True)
    if monster_type == "pvp":
        weaken_uptime = min(eff_duration(30), PVP_FIGHT_DURATION) / PVP_FIGHT_DURATION
    else:
        weaken_uptime = eff_duration(30) / fb_eff_cd
    monster_dmg_bonus = 15 * weaken_uptime
else:
    monster_dmg_bonus = 0.0
monster_dmg_bonus += elemental_reset_pct

# Frozen Break (Lv.66 base, +2%p/stack Mastery Lv.98, always 5 stacks) + Frost Clutch (Lv.125,
# 4%/stack patched, always 5 stacks, applies to all skills) — flat, level-gated additive Damage%.
damage_bonus = (15 if level >= 66 else 0) + (10 if level >= 98 else 0) + (20 if level >= 125 else 0)

# Average buff multiplier. Magic Guard and Meditation are both "+X% Attack" sources — same
# bucket, so they sum into one combined percentage before a single multiplication; Infinity is
# Final Damage, a different (and deliberately still multiplicative) bucket, so it stays its own
# separate factor.
attack_bucket_sum = 0.0
final_damage_mult = 1.0
for key, b in BUFFS.items():
    if not unlocked(key):
        continue
    pct = coeff_pct(b["base"], b["fidx"], b["scales"], b["job_step"])
    if "ramp_mult" in b:
        pct *= b["ramp_mult"]
    eff_cd = eff_cooldown(b["cooldown"], b["costs_action"])
    if fixed_duration_active:
        uptime_fraction = exact_buff_uptime(eff_cd, eff_duration(b["duration"])) / fight_duration
    elif monster_type == "pvp":
        uptime_fraction = min(eff_duration(b["duration"]), PVP_FIGHT_DURATION) / PVP_FIGHT_DURATION
    else:
        uptime_fraction = eff_duration(b["duration"]) / eff_cd
    avg_pct = pct * uptime_fraction
    if b["target"] == "ATTACK":
        attack_bucket_sum += avg_pct
    else:
        final_damage_mult *= 1 + avg_pct / 100
avg_buff_mult = (1 + attack_bucket_sum / 100) * final_damage_mult

# MP Eater - MP Boost: flat +7% Attack Speed once unlocked, no uptime averaging (not a timed buff).
as_bonus = mp_eater_as_pct if unlocked("MP_EATER_MP_BOOST") else 0.0

actions_per_second = 1 + min(150, 150 * (1 - (1 - attack_speed_base / 150) * (1 - as_bonus / 150))) / 100

# Cooldown lookup used by the cast-rate sum below — Blizzard - Final Attack shares Blizzard's
# own (CDR-eligible) cooldown, not its own CostsActionSlot flag (see build script's own
# CDR_COSTS_ACTION_ROW override and comment for why).
_CDR_COSTS_ACTION_OVERRIDE = {"BLIZZARD_FINAL_ATTACK": "BLIZZARD"}


def _cdr_costs_action(key, s):
    ref_key = _CDR_COSTS_ACTION_OVERRIDE.get(key, key)
    return SKILLS[ref_key]["costs_action"]


# Buff-Casting Startup Delay: in fixed-duration fights, the character casts every currently-
# unlocked, actively-cast buff sequentially at fight start (1/APS seconds each, same cadence as
# every other action-costing skill) before their first damage-skill cast — so damage skills'
# usable window is reduced by this amount. Buffs themselves keep their own t=0 uptime math
# unchanged (they're what causes the delay, not affected further by it).
if fixed_duration_active:
    buff_cast_startup_time = sum(1 for key, b in BUFFS.items() if b["costs_action"] and unlocked(key)) / actions_per_second
else:
    buff_cast_startup_time = 0.0


def non_buff_casts(cooldown):
    """CastsInFight for a non-buff (damage) skill, reduced by the buff-casting startup delay."""
    if buff_cast_startup_time >= fight_duration:
        return 0
    return exact_casts(cooldown, max(0.0, fight_duration - buff_cast_startup_time))


def non_buff_total_hits(cooldown, hits_per_cast, icd, window):
    """exact_total_hits for a non-buff (damage) skill, reduced by the startup delay."""
    if buff_cast_startup_time >= fight_duration:
        return 0.0
    return exact_total_hits(cooldown, hits_per_cast, icd, window, max(0.0, fight_duration - buff_cast_startup_time))


if fixed_duration_active:
    cast_rate = sum(
        non_buff_casts(eff_cooldown(s["cooldown"], _cdr_costs_action(k, s)))
        for k, s in SKILLS.items() if s["costs_action"] and unlocked(k)
    )
    cast_rate += sum(
        exact_casts(eff_cooldown(b["cooldown"], b["costs_action"]))
        for key, b in BUFFS.items() if b["costs_action"] and unlocked(key)
    )
    cast_rate /= fight_duration
else:
    cast_rate = sum(
        1 / eff_cooldown(s["cooldown"], _cdr_costs_action(k, s))
        for k, s in SKILLS.items() if s["costs_action"] and unlocked(k)
    )
    cast_rate += sum(
        1 / eff_cooldown(b["cooldown"], b["costs_action"])
        for key, b in BUFFS.items() if b["costs_action"] and unlocked(key)
    )
chain_lightning_per_second = max(0, actions_per_second - cast_rate)

maple_lvl = input_level(4, level, skill1, skill2, skill3, skill4, skill_all)
maple_factor = get_factor(maple_lvl, 23)


def maple_mult(key):
    if key not in MAPLE_TARGETS:
        return 1.0
    return 1 + (MAPLE_TARGETS[key] / 10) * (maple_factor / 1000) / 100


def hit_damage(coeff_pct_val, is_basic, maple_mult_val, s):
    effective_coeff = coeff_pct_val + s["mastery"]
    base_damage = attack * (effective_coeff / 100)
    dmg_reduction = 5000 / (6000 + monster_defense * (1 - def_pen / 100))
    elem_amp = element_amp_pct if unlocked("ELEMENT_AMPLIFICATION") else 0.0
    arcane = arcane_aim_pct if unlocked("ARCANE_AIM") else 0.0
    final_mult = (1 + final_damage / 100) * (1 + elem_amp / 100) * (1 + arcane / 100) ** 5
    source_pct = basic_attack_damage if is_basic else skill_damage
    # Boss/Normal Monster Damage% is deliberately NOT applied here — it's blended per-branch in
    # boss_normal_multiplier (or FROZEN_ORB's own inline split below), applied by the caller.
    extra_mult = avg_buff_mult * maple_mult_val
    base_hit = (
        base_damage * (1 + stat_damage / 100) * (1 + (damage + damage_bonus) / 100)
        * (1 + damage_amp / 100) * dmg_reduction * final_mult * (1 + source_pct / 100) * extra_mult
    )
    non_crit_min = base_hit * min(min_damage, max_damage) / 100
    non_crit_max = base_hit * max_damage / 100
    non_crit_avg = (non_crit_min + non_crit_max) / 2
    crit_avg = non_crit_avg * (1 + crit_damage / 100)
    cr = min(crit_rate, 100) / 100
    return non_crit_avg * (1 - cr) + crit_avg * cr


chain_s = SKILLS["CHAIN_LIGHTNING"]
chain_hit = hit_damage(skill_coefficient_base, True, 1.0, chain_s)
chain_dps = chain_lightning_hits * chain_hit * chain_lightning_per_second * boss_normal_multiplier("CHAIN_LIGHTNING", chain_s) \
    if unlocked("CHAIN_LIGHTNING") else 0.0

skill_dps = {}
for key, s in SKILLS.items():
    if key == "CHAIN_LIGHTNING":
        continue
    if not unlocked(key):
        skill_dps[key] = 0.0
        continue
    pct = coeff_pct(s["base"], s["fidx"], s["scales"], s["job_step"])
    proc = 1 - (1 - s.get("chance", 100) / 100) ** s.get("rolls", 1)
    hd = hit_damage(pct, False, maple_mult(key), s)
    eff_cd = eff_cooldown(s["cooldown"], _cdr_costs_action(key, s))
    if fixed_duration_active:
        rate = non_buff_total_hits(eff_cd, s["hits"], s.get("icd"), s.get("window")) / fight_duration
    else:
        hits = s["hits"] * ((s["window"] / s["icd"]) if s.get("icd") else 1)
        rate = hits / eff_cd
    if key == "FROZEN_ORB":
        # Frozen Orb's own boss/normal damage multiplier (0.5 vs single target, 1.0 vs
        # normal-monster AoE) is applied INSIDE each branch, not as a shared w-blended scalar —
        # see build_ice_lightning_mage_workbook.py's own comment on this branch for why.
        if monster_type == "pvp":
            boss_mult, normal_mult = 1.0, 1.0
        else:
            boss_term = boss_damage + s.get("mastery_boss", 0) + monster_dmg_bonus
            normal_term = normal_damage + s.get("mastery_normal", 0) + monster_dmg_bonus
            targets = min(NORMAL_MONSTER_TARGETS.get(key, 1), max_enemies_hit)
            boss_mult = 1 + boss_term / 100
            normal_mult = (1 + normal_term / 100) * targets
        combined_mult = (1 - normal_weight) * 0.5 * boss_mult + normal_weight * 1.0 * normal_mult
        skill_dps[key] = rate * proc * hd * combined_mult
    else:
        skill_dps[key] = rate * proc * hd * boss_normal_multiplier(key, s)

total_dps = chain_dps + sum(skill_dps.values())

print(f"Monster Damage Taken bonus % (Freezing Breath Weaken + Elemental Reset) = {monster_dmg_bonus:.6f}")
print(f"Damage % bonus (Frozen Break + Frost Clutch) = {damage_bonus:.6f}")
print(f"Average buff multiplier    = {avg_buff_mult:.6f}")
print(f"Attack speed bonus % (MP Eater) = {as_bonus:.6f}")
print(f"Actions/sec                = {actions_per_second:.6f}")
print(f"Cast rate (subtracted)     = {cast_rate:.6f}")
print(f"Chain Lightning casts/sec  = {chain_lightning_per_second:.6f}")
print(f"Chain Lightning DPS        = {chain_dps:.4f}")
for k, v in skill_dps.items():
    print(f"  {k:28s} DPS = {v:.4f}")
print(f"TOTAL DPS                  = {total_dps:.4f}")

# ---------------------------------------------------------------------------
# Cross-check against the actual live workbook (evaluated with `formulas`)
# ---------------------------------------------------------------------------
import formulas  # noqa: E402

print()
print(f"Loading live workbook: {XLSX_PATH}")
xl_model = formulas.ExcelModel().loads(str(XLSX_PATH)).finish()
solution = xl_model.calculate()


def scalar(value):
    return float(np.asarray(value).reshape(-1)[0])


def read_cell(sheet, cell):
    suffix = f"{sheet.upper()}'!{cell.upper()}"
    for k, v in solution.items():
        if k.upper().endswith(suffix):
            return scalar(v.value)
    raise KeyError(f"cell not found in workbook: {sheet}!{cell}")


py_skill_dps = dict(skill_dps)
py_skill_dps["CHAIN_LIGHTNING"] = chain_dps

xl_total = read_cell("SUMMARY", f"B{SUMMARY_ROW['TOTAL_DPS']}")
mismatches = []
print()
print(f"{'Skill':28s} {'Python DPS':>16s} {'Workbook DPS':>16s}  match")
for key, row in ROW.items():
    py_val = py_skill_dps.get(key, 0.0)
    xl_val = read_cell("CALC", f"O{row}")
    ok = np.isclose(py_val, xl_val, rtol=1e-6, atol=1e-3)
    print(f"{key:28s} {py_val:16.4f} {xl_val:16.4f}  {'OK' if ok else 'MISMATCH'}")
    if not ok:
        mismatches.append((key, py_val, xl_val))

print()
total_ok = np.isclose(total_dps, xl_total, rtol=1e-6, atol=1e-3)
print(f"{'TOTAL DPS':28s} {total_dps:16.4f} {xl_total:16.4f}  {'OK' if total_ok else 'MISMATCH'}")
if not total_ok:
    mismatches.append(("TOTAL DPS", total_dps, xl_total))

if mismatches:
    print(f"\n{len(mismatches)} mismatch(es) between the independent model and the live workbook.")
    sys.exit(1)
print("\nAll rows match between the independent model and the live workbook.")
